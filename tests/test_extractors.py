from io import BytesIO

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from m365_mcp.extractors import (
    AttachmentExtractorRegistry,
    AttachmentInput,
    AttachmentTooLargeError,
    UnsupportedAttachmentError,
)


def test_builtin_extractors_read_supported_formats() -> None:
    registry = AttachmentExtractorRegistry()

    text = registry.extract(
        AttachmentInput(
            name="notes.txt",
            content_type="text/plain; charset=utf-8",
            content="plain text".encode(),
        )
    )
    assert text.format == "txt"
    assert text.content == "plain text"

    pdf = registry.extract(
        AttachmentInput(
            name="report.pdf",
            content_type="application/pdf",
            content=_pdf_with_text("PDF body"),
        )
    )
    assert pdf.format == "pdf"
    assert "PDF body" in pdf.content

    document = Document()
    document.add_paragraph("DOCX body")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "left"
    table.cell(0, 1).text = "right"
    docx_bytes = BytesIO()
    document.save(docx_bytes)
    docx = registry.extract(
        AttachmentInput(
            name="report.docx",
            content_type=(
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            ),
            content=docx_bytes.getvalue(),
        )
    )
    assert docx.format == "docx"
    assert "DOCX body" in docx.content
    assert "left\tright" in docx.content

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Data"
    worksheet.append(["name", "value"])
    worksheet.append(["alpha", 42])
    xlsx_bytes = BytesIO()
    workbook.save(xlsx_bytes)
    workbook.close()
    xlsx = registry.extract(
        AttachmentInput(
            name="report.xlsx",
            content_type=(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            ),
            content=xlsx_bytes.getvalue(),
        )
    )
    assert xlsx.format == "xlsx"
    assert "[Sheet Data]" in xlsx.content
    assert "alpha\t42" in xlsx.content

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    textbox.text = "PPTX body"
    pptx_bytes = BytesIO()
    presentation.save(pptx_bytes)
    pptx = registry.extract(
        AttachmentInput(
            name="slides.pptx",
            content_type=(
                "application/vnd.openxmlformats-officedocument.presentationml.presentation"
            ),
            content=pptx_bytes.getvalue(),
        )
    )
    assert pptx.format == "pptx"
    assert "PPTX body" in pptx.content


def test_extractor_enforces_byte_and_text_limits() -> None:
    registry = AttachmentExtractorRegistry(max_bytes=4, max_text_chars=10)

    with pytest.raises(AttachmentTooLargeError):
        registry.extract(AttachmentInput(name="notes.txt", content=b"12345"))

    with pytest.raises(UnsupportedAttachmentError):
        AttachmentExtractorRegistry().extract(
            AttachmentInput(name="archive.zip", content=b"PK")
        )

    result = AttachmentExtractorRegistry(max_text_chars=5).extract(
        AttachmentInput(name="notes.txt", content=b"0123456789")
    )
    assert result.truncated is True
    assert len(result.content) == 5


def _pdf_with_text(text: str) -> bytes:
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>"
        ),
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    document = b"%PDF-1.4\n"
    offsets = [0]
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(document))
        document += f"{number} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(document)
    document += f"xref\n0 {len(objects) + 1}\n".encode()
    document += b"0000000000 65535 f \n"
    for offset in offsets[1:]:
        document += f"{offset:010d} 00000 n \n".encode()
    document += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref}\n%%EOF\n"
    ).encode()
    return document
