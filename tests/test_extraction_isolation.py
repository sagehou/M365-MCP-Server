import asyncio
import subprocess
import sys
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest

from m365_mcp.extractors import (
    AttachmentExtractorRegistry, AttachmentInput, AttachmentTooLargeError, InvalidAttachmentError,
)
from m365_mcp.extractors import isolation


def test_office_archive_expanded_limit_rejects_small_zip_bomb():
    stream = BytesIO()
    with ZipFile(stream, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"x" * (65 * 1024 * 1024))
    assert len(stream.getvalue()) < 100_000
    with pytest.raises(AttachmentTooLargeError):
        AttachmentExtractorRegistry().extract(AttachmentInput(name="bomb.docx", content=stream.getvalue()))


def test_real_worker_reads_text_and_rejects_corrupt_pdf():
    async def exercise():
        registry = AttachmentExtractorRegistry(max_text_chars=5)
        result = await registry.extract_async(AttachmentInput(name="notes.txt", content=b"0123456789"))
        assert result.content == "01234"
        assert result.truncated
        with pytest.raises(InvalidAttachmentError):
            await registry.extract_async(AttachmentInput(name="bad.pdf", content=b"not a pdf"))
    asyncio.run(exercise())


@pytest.mark.parametrize("cancel", [False, True])
def test_timeout_and_cancellation_kill_and_reap_worker(monkeypatch, cancel):
    class HungWorker:
        returncode = None
        killed = False
        reaped = False
        async def communicate(self, data):
            await asyncio.Event().wait()
        def kill(self):
            self.killed = True
            self.returncode = -9
        async def wait(self):
            self.reaped = True
            return self.returncode
    worker = HungWorker()
    async def spawn(*args, **kwargs):
        assert "CLIENT_SECRET" not in kwargs["env"]
        return worker
    monkeypatch.setenv("CLIENT_SECRET", "DO_NOT_INHERIT")
    monkeypatch.setattr(asyncio, "create_subprocess_exec", spawn)
    monkeypatch.setattr(isolation, "WORKER_TIMEOUT_SECONDS", 0.01)
    async def exercise():
        task = asyncio.create_task(isolation.extract_isolated(
            AttachmentInput(name="notes.txt", content=b"hello"), max_bytes=10, max_text_chars=10))
        if cancel:
            await asyncio.sleep(0)
            task.cancel()
        with pytest.raises(asyncio.CancelledError if cancel else InvalidAttachmentError):
            await task
    asyncio.run(exercise())
    assert worker.killed and worker.reaped


def test_worker_installs_hard_resource_limits(monkeypatch):
    from m365_mcp.extractors import worker
    calls = {}
    monkeypatch.setattr(worker.resource, "setrlimit", lambda key, value: calls.update({key: value}))
    worker.apply_limits()
    assert calls[worker.resource.RLIMIT_AS] == (512 * 1024 * 1024,) * 2
    assert calls[worker.resource.RLIMIT_CPU] == (15, 15)
    assert calls[worker.resource.RLIMIT_CORE] == (0, 0)


def test_real_worker_address_space_limit_rejects_oversized_allocation():
    result = subprocess.run([
        sys.executable, "-c",
        "from m365_mcp.extractors.worker import apply_limits; "
        "apply_limits(); bytearray(600 * 1024 * 1024)",
    ], capture_output=True, timeout=10)
    assert result.returncode != 0
    assert b"MemoryError" in result.stderr


def test_isolated_worker_supports_all_document_formats():
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from test_extractors import _pdf_with_text

    document = Document()
    document.add_paragraph("DOCX body")
    docx = BytesIO()
    document.save(docx)
    workbook = Workbook()
    workbook.active.append(["XLSX body"])
    xlsx = BytesIO()
    workbook.save(xlsx)
    workbook.close()
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = "PPTX body"
    pptx = BytesIO()
    presentation.save(pptx)

    async def exercise():
        registry = AttachmentExtractorRegistry()
        for name, content, expected in [
            ("notes.txt", b"TXT body", "TXT body"),
            ("notes.pdf", _pdf_with_text("PDF body"), "PDF body"),
            ("notes.docx", docx.getvalue(), "DOCX body"),
            ("notes.xlsx", xlsx.getvalue(), "XLSX body"),
            ("notes.pptx", pptx.getvalue(), "PPTX body"),
        ]:
            result = await registry.extract_async(AttachmentInput(name=name, content=content))
            assert expected in result.content
    asyncio.run(exercise())


def test_registry_limits_active_parser_workers(monkeypatch):
    from m365_mcp.extractors import ExtractionResult
    active = maximum = 0
    async def extract(*args, **kwargs):
        nonlocal active, maximum
        active += 1
        maximum = max(maximum, active)
        await asyncio.sleep(0.01)
        active -= 1
        return ExtractionResult("txt", "ok", False)
    monkeypatch.setattr(isolation, "extract_isolated", extract)
    async def exercise():
        registry = AttachmentExtractorRegistry()
        await asyncio.gather(*[
            registry.extract_async(AttachmentInput(name="notes.txt", content=b"ok"))
            for _ in range(6)
        ])
    asyncio.run(exercise())
    assert maximum == 2
