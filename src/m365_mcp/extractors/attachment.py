"""Bounded, server-side text extraction for supported file attachments."""

from __future__ import annotations

import asyncio
from collections.abc import Iterable
from dataclasses import dataclass
from io import BytesIO
from typing import ClassVar, Protocol
from zipfile import BadZipFile, ZipFile

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

from .errors import (
    AttachmentTooLargeError,
    InvalidAttachmentError,
    UnsupportedAttachmentError,
)


DEFAULT_MAX_BYTES = 10 * 1024 * 1024
DEFAULT_MAX_TEXT_CHARS = 100_000
TRUNCATION_MARKER = "\n\n[Content truncated by server limit.]"


@dataclass(frozen=True, slots=True)
class AttachmentInput:
    """Attachment metadata and bytes obtained from Microsoft Graph."""

    name: str
    content: bytes
    content_type: str | None = None
    declared_size: int | None = None


@dataclass(frozen=True, slots=True)
class ExtractionResult:
    """Bounded text plus metadata suitable for an MCP response."""

    format: str
    content: str
    truncated: bool


class AttachmentExtractor(Protocol):
    """Protocol implemented by each supported document format handler."""

    format_name: str
    extensions: ClassVar[frozenset[str]]
    media_types: ClassVar[frozenset[str]]

    def extract(self, content: bytes) -> str:
        """Return text extracted from attachment bytes."""


class TextAttachmentExtractor:
    format_name = "txt"
    extensions = frozenset({".txt"})
    media_types = frozenset({"text/plain"})

    def extract(self, content: bytes) -> str:
        return content.decode("utf-8-sig", errors="replace")


class MarkdownAttachmentExtractor:
    format_name = "markdown"
    extensions = frozenset({".md", ".markdown"})
    media_types = frozenset({"text/markdown", "text/x-markdown"})

    def extract(self, content: bytes) -> str:
        return content.decode("utf-8-sig", errors="replace")


class PdfAttachmentExtractor:
    format_name = "pdf"
    extensions = frozenset({".pdf"})
    media_types = frozenset({"application/pdf"})

    def extract(self, content: bytes) -> str:
        try:
            reader = PdfReader(BytesIO(content), strict=False)
            pages: list[str] = []
            for number, page in enumerate(reader.pages, start=1):
                text = page.extract_text() or ""
                if text.strip():
                    pages.append(f"[Page {number}]\n{text.strip()}")
            return "\n\n".join(pages)
        except Exception as exc:
            raise InvalidAttachmentError("The PDF attachment could not be parsed") from exc


class DocxAttachmentExtractor:
    format_name = "docx"
    extensions = frozenset({".docx"})
    media_types = frozenset(
        {
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        }
    )

    def extract(self, content: bytes) -> str:
        try:
            document = Document(BytesIO(content))
            parts = [
                paragraph.text.strip()
                for paragraph in document.paragraphs
                if paragraph.text.strip()
            ]
            for number, table in enumerate(document.tables, start=1):
                rows = []
                for row in table.rows:
                    cells = [cell.text.replace("\n", " ").strip() for cell in row.cells]
                    if any(cells):
                        rows.append("\t".join(cells))
                if rows:
                    parts.append(f"[Table {number}]\n" + "\n".join(rows))
            return "\n\n".join(parts)
        except Exception as exc:
            raise InvalidAttachmentError("The DOCX attachment could not be parsed") from exc


class XlsxAttachmentExtractor:
    format_name = "xlsx"
    extensions = frozenset({".xlsx"})
    media_types = frozenset(
        {
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        }
    )

    def extract(self, content: bytes) -> str:
        workbook = None
        try:
            workbook = load_workbook(
                filename=BytesIO(content),
                read_only=True,
                data_only=True,
            )
            sheets: list[str] = []
            for worksheet in workbook.worksheets:
                rows: list[str] = []
                for row in worksheet.iter_rows(values_only=True):
                    values = ["" if value is None else str(value) for value in row]
                    if any(values):
                        rows.append("\t".join(values).rstrip())
                if rows:
                    sheets.append(f"[Sheet {worksheet.title}]\n" + "\n".join(rows))
            return "\n\n".join(sheets)
        except Exception as exc:
            raise InvalidAttachmentError("The XLSX attachment could not be parsed") from exc
        finally:
            if workbook is not None:
                workbook.close()


class PptxAttachmentExtractor:
    format_name = "pptx"
    extensions = frozenset({".pptx"})
    media_types = frozenset(
        {
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        }
    )

    def extract(self, content: bytes) -> str:
        try:
            presentation = Presentation(BytesIO(content))
            slides: list[str] = []
            for number, slide in enumerate(presentation.slides, start=1):
                elements: list[str] = []
                for shape in slide.shapes:
                    if getattr(shape, "has_table", False):
                        rows = []
                        for row in shape.table.rows:
                            cells = [
                                cell.text.replace("\n", " ").strip()
                                for cell in row.cells
                            ]
                            if any(cells):
                                rows.append("\t".join(cells))
                        if rows:
                            elements.append("\n".join(rows))
                    elif getattr(shape, "has_text_frame", False):
                        text = shape.text.strip()
                        if text:
                            elements.append(text)
                if elements:
                    slides.append(f"[Slide {number}]\n" + "\n".join(elements))
            return "\n\n".join(slides)
        except Exception as exc:
            raise InvalidAttachmentError("The PPTX attachment could not be parsed") from exc


class AttachmentExtractorRegistry:
    """Resolve an attachment format and enforce extraction output limits."""

    def __init__(
        self,
        extractors: Iterable[AttachmentExtractor] | None = None,
        *,
        max_bytes: int = DEFAULT_MAX_BYTES,
        max_text_chars: int = DEFAULT_MAX_TEXT_CHARS,
    ) -> None:
        if max_bytes <= 0:
            raise ValueError("max_bytes must be positive")
        if max_text_chars <= 0:
            raise ValueError("max_text_chars must be positive")
        self.max_bytes = max_bytes
        self.max_text_chars = max_text_chars
        self._slots = asyncio.Semaphore(2)
        self._custom_handlers = extractors is not None
        handlers = tuple(extractors or default_extractors())
        self._by_extension: dict[str, AttachmentExtractor] = {}
        self._by_media_type: dict[str, AttachmentExtractor] = {}
        for handler in handlers:
            for extension in handler.extensions:
                self._register(self._by_extension, extension.casefold(), handler)
            for media_type in handler.media_types:
                self._register(self._by_media_type, media_type.casefold(), handler)

    @staticmethod
    def _register(
        target: dict[str, AttachmentExtractor],
        key: str,
        handler: AttachmentExtractor,
    ) -> None:
        if key in target:
            raise ValueError(f"Duplicate attachment extractor registration: {key}")
        target[key] = handler

    def extract(self, attachment: AttachmentInput) -> ExtractionResult:
        """Extract bounded text without exposing the source bytes."""

        if attachment.declared_size is not None and attachment.declared_size > self.max_bytes:
            raise AttachmentTooLargeError(
                f"Attachment exceeds the {self.max_bytes}-byte server limit"
            )
        if len(attachment.content) > self.max_bytes:
            raise AttachmentTooLargeError(
                f"Attachment exceeds the {self.max_bytes}-byte server limit"
            )

        handler = self._resolve(attachment)
        if handler.format_name in {"docx", "xlsx", "pptx"}:
            _validate_archive(attachment.content)
        try:
            extracted = handler.extract(attachment.content)
        except (AttachmentTooLargeError, InvalidAttachmentError):
            raise
        except Exception as exc:
            raise InvalidAttachmentError("The attachment could not be parsed") from exc

        bounded, truncated = _truncate(extracted, self.max_text_chars)
        return ExtractionResult(
            format=handler.format_name,
            content=bounded,
            truncated=truncated,
        )

    async def extract_async(self, attachment: AttachmentInput) -> ExtractionResult:
        """Run built-in parsers outside the web process with a hard deadline."""
        from .isolation import extract_isolated

        if self._custom_handlers:
            raise ValueError("Isolated extraction requires registered built-in handlers")
        async with self._slots:
            return await extract_isolated(
                attachment, max_bytes=self.max_bytes, max_text_chars=self.max_text_chars
            )

    def _resolve(self, attachment: AttachmentInput) -> AttachmentExtractor:
        filename = attachment.name.replace("\\", "/").rsplit("/", 1)[-1]
        if not filename:
            raise InvalidAttachmentError("The attachment has no file name")
        extension = ""
        if "." in filename and not filename.endswith("."):
            extension = "." + filename.rsplit(".", 1)[1].casefold()
        media_type = (attachment.content_type or "").split(";", 1)[0].strip().casefold()
        handler = self._by_extension.get(extension) or self._by_media_type.get(media_type)
        if handler is None:
            raise UnsupportedAttachmentError(
                "Attachment format is not supported; use PDF, DOCX, XLSX, PPTX, TXT, or Markdown"
            )
        return handler


def _validate_archive(content: bytes) -> None:
    """Reject expanded Office archives before document libraries load XML."""
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 2048 or sum(item.file_size for item in entries) > 64 * 1024 * 1024:
                raise AttachmentTooLargeError("Office archive exceeds expanded size or member limit")
            if any(item.flag_bits & 1 for item in entries):
                raise InvalidAttachmentError("Encrypted Office archives are not supported")
    except BadZipFile:
        raise InvalidAttachmentError("The Office attachment is not a valid archive") from None


def _truncate(text: str, max_chars: int) -> tuple[str, bool]:
    if len(text) <= max_chars:
        return text, False
    if max_chars <= len(TRUNCATION_MARKER):
        return text[:max_chars], True
    return text[: max_chars - len(TRUNCATION_MARKER)] + TRUNCATION_MARKER, True


def default_extractors() -> tuple[AttachmentExtractor, ...]:
    """Return the built-in handlers in deterministic registration order."""

    return (
        TextAttachmentExtractor(),
        MarkdownAttachmentExtractor(),
        PdfAttachmentExtractor(),
        DocxAttachmentExtractor(),
        XlsxAttachmentExtractor(),
        PptxAttachmentExtractor(),
    )
